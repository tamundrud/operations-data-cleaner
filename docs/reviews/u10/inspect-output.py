"""Independent check: python3 docs/reviews/u10/inspect-output.py (openpyxl 3.1.5)."""
from pathlib import Path
import json
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook, __version__

path = Path(__file__).with_name("synthetic-output.xlsx")
workbook = load_workbook(path, data_only=False)
assert workbook.sheetnames == ["Output"]
sheet = workbook["Output"]
assert sheet["A2"].value == "001234" and sheet["A2"].data_type == "s"
assert sheet["B2"].value == 0.3 and sheet["B2"].data_type == "n"
assert sheet["C2"].value.isoformat() == "2024-02-29T00:00:00"
assert sheet["C2"].number_format == "mm/dd/yyyy"
assert sheet["D2"].value is None
assert sheet["E2"].value == '=HYPERLINK("literal")'
assert sheet["E2"].data_type == "s"
with zipfile.ZipFile(path) as archive:
    for name in archive.namelist():
        if name.endswith((".xml", ".rels")):
            root = ET.fromstring(archive.read(name))
            assert not any(
                node.tag.rsplit("}", 1)[-1] in ("f", "hyperlink")
                or node.attrib.get("TargetMode") == "External"
                for node in root.iter()
            )
print(json.dumps({"reader": "openpyxl", "version": __version__, "status": "passed"}))
